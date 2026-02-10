from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

REQUIRED_COLS = {
    "НомерПартии": "batch_number",
    "ДатаПартии": "batch_date",
    "Номенклатура": "nomenclature",
    "РабочийЦентр": "work_center_name",
    "ИдентификаторРЦ": "work_center_identifier",
    "Смена": "shift",
    "Бригада": "team",
    "ПредставлениеЗаданияНаСмену": "task_description",
    "ДатаВремяНачалаСмены": "shift_start",
    "ДатаВремяОкончанияСмены": "shift_end",
    "КодЕКН": "ekn_code",
    "СтатусЗакрытия": "is_closed",
}

def parse_batches_excel(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active


    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(header) if name}

    missing = [col for col in REQUIRED_COLS.keys() if col not in idx]
    if missing:
        raise ValueError(f"Missing columns: {missing}")

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        item = {}
        for col_ru, key in REQUIRED_COLS.items():
            item[key] = row[idx[col_ru]]


        if isinstance(item["batch_date"], datetime):
            item["batch_date"] = item["batch_date"].date()
        if isinstance(item["shift_start"], str):
            item["shift_start"] = datetime.fromisoformat(item["shift_start"])
        if isinstance(item["shift_end"], str):
            item["shift_end"] = datetime.fromisoformat(item["shift_end"])

        rows.append(item)

    return rows
